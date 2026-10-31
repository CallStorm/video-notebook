#!/usr/bin/env python3
"""多格式文档文本提取器（思维导图生成技能用）。

支持：
  - .epub : 标准库 zipfile 解析，零依赖
  - .txt/.md : 直接读取
  - .pdf  : 需 pdfplumber 或 pdfminer.six（缺失会提示）
  - .docx : 需 python-docx（缺失会提示）

用法：
  python extract_text.py <输入文件> [--out 输出.txt]
不指定 --out 时直接打印到 stdout。
"""
import sys
import os
import re
import html
import zipfile
import argparse


# 视为非内容页、提取时跳过的章节关键词（大小写不敏感）
SKIP_KEYWORDS = ('titlepage', 'cover', 'copyright', 'toc', 'nav')


def extract_epub(path):
    z = zipfile.ZipFile(path)
    # 兼容不同打包结构：章节可能在 text/、OEBPS/xhtml/ 等任意位置
    def sort_key(n):
        nums = re.findall(r'\d+', n)
        return [int(x) for x in nums] if nums else [0]

    names = [
        n for n in z.namelist()
        if re.search(r'\.x?html?$', n, re.I)
        and not any(k in n.lower() for k in SKIP_KEYWORDS)
    ]
    names = sorted(names, key=sort_key)
    out = []
    for n in names:
        try:
            data = z.read(n).decode('utf-8', 'ignore')
        except Exception:
            continue
        data = re.sub(r'<script.*?</script>', '', data, flags=re.S | re.I)
        data = re.sub(r'<style.*?</style>', '', data, flags=re.S | re.I)
        text = re.sub(r'<[^>]+>', ' ', data)
        text = html.unescape(text)
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n\s*\n+', '\n', text).strip()
        if text:
            out.append(text)
    return '\n\n'.join(out)


def extract_plain(path):
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()


def extract_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        try:
            from pdfminer.high_level import extract_text as pm_extract
            return pm_extract(path)
        except ImportError:
            raise RuntimeError(
                "PDF 提取需要 pdfplumber 或 pdfminer.six，请安装：pip install pdfplumber"
            )
    out = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ''
            out.append(t)
    return '\n'.join(out)


def extract_docx(path):
    try:
        import docx
    except ImportError:
        raise RuntimeError("DOCX 提取需要 python-docx，请安装：pip install python-docx")
    d = docx.Document(path)
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            paras.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(paras)


def main():
    ap = argparse.ArgumentParser(description="提取文档纯文本")
    ap.add_argument('input', help='输入文件路径')
    ap.add_argument('--out', help='输出文本文件路径（不指定则打印到 stdout）')
    args = ap.parse_args()

    ext = os.path.splitext(args.input)[1].lower()
    try:
        if ext == '.epub':
            text = extract_epub(args.input)
        elif ext in ('.txt', '.md', '.markdown'):
            text = extract_plain(args.input)
        elif ext == '.pdf':
            text = extract_pdf(args.input)
        elif ext in ('.docx', '.doc'):
            text = extract_docx(args.input)
        else:
            text = extract_plain(args.input)
    except RuntimeError as e:
        sys.stderr.write(f"[错误] {e}\n")
        sys.exit(2)

    if args.out:
        with open(args.out, 'w', encoding='utf-8') as f:
            f.write(text)
        sys.stderr.write(f"已写出 {len(text)} 字符到 {args.out}\n")
    else:
        sys.stdout.write(text)


if __name__ == '__main__':
    main()
